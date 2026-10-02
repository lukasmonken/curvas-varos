import openpyxl, sys
from openpyxl.worksheet.formula import ArrayFormula
wbf=openpyxl.load_workbook('docs/legado/Query-Avila.xlsx',data_only=False)
wbv=openpyxl.load_workbook('docs/legado/Query-Avila.xlsx',data_only=True)
name=sys.argv[1]; maxr=int(sys.argv[2]); 
wsf=wbf[name]; wsv=wbv[name]
for row in wsf.iter_rows(min_row=1,max_row=maxr):
    for c in row:
        if c.value is None: continue
        v=wsv[c.coordinate].value
        f=c.value
        if isinstance(f,ArrayFormula): f='{ARRAY '+f.ref+'} '+str(f.text)
        fill=c.fill.fgColor.rgb if c.fill and c.fill.fill_type else ''
        print(f"{c.coordinate}\t{f!r}\t=> {v!r}\t{fill}")
