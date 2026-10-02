import openpyxl, re, sys
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.utils import get_column_letter, column_index_from_string
wb=openpyxl.load_workbook('docs/legado/Query-Avila.xlsx',data_only=False)
wv=openpyxl.load_workbook('docs/legado/Query-Avila.xlsx',data_only=True)
ref=re.compile(r"(\$?)([A-Z]{1,3})(\$?)(\d+)")
def rel(f,r,c):
    def sub(m):
        cd,col,rd,row=m.groups()
        ci=column_index_from_string(col); ri=int(row)
        cs=f"C{ci}" if cd else f"C[{ci-c}]"
        rs=f"R{ri}" if rd else f"R[{ri-r}]"
        return rs+cs
    return ref.sub(sub,f)
for name in sys.argv[1:]:
    ws=wb[name]; vs=wv[name]
    print("=====",name)
    for col in range(1,ws.max_column+1):
        runs=[]; 
        for r in range(1,ws.max_row+1):
            v=ws.cell(r,col).value
            if v is None: continue
            if isinstance(v,ArrayFormula): v='{A}'+str(v.text)
            key=rel(v,r,col) if isinstance(v,str) and v.startswith(('=','{A}')) else ('<const>' if not isinstance(v,str) else '<text>')
            if runs and runs[-1][0]==key and runs[-1][2]==r-1: runs[-1][2]=r
            else: runs.append([key,r,r])
        L=get_column_letter(col)
        for k,a,b in runs:
            if b-a>=5: print(f"{L}{a}:{L}{b}  ({b-a+1}) {k}  first={ws.cell(a,col).value!r} lastval={vs.cell(b,col).value!r}")
