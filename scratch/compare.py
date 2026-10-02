import openpyxl, math
orig_f=openpyxl.load_workbook('docs/legado/Query-Avila.xlsx',data_only=False)
orig=openpyxl.load_workbook('docs/legado/Query-Avila.xlsx',data_only=True)
rc=openpyxl.load_workbook('scratch/recalc/QUERY-AVILA.XLSX',data_only=True)
smap={n.upper():n for n in orig.sheetnames}
rcs={s.title.upper():s for s in rc.worksheets}
print(rc.sheetnames)
worst={}
cnt=0
for name in orig.sheetnames:
    wf=orig_f[name]; wo=orig[name]; wr=rcs[name.upper()]
    mx=(0,None)
    for row in wf.iter_rows():
        for c in row:
            v=c.value
            if not (isinstance(v,str) and v.startswith('=')) and type(v).__name__!='ArrayFormula': continue
            a=wo[c.coordinate].value; b=wr[c.coordinate].value
            cnt+=1
            if isinstance(a,(int,float)) and isinstance(b,(int,float)):
                d=abs(a-b)
                if d>mx[0]: mx=(d,c.coordinate,a,b)
            elif a!=b and not (a is None and b in ('',None)):
                print('NONNUM DIFF',name,c.coordinate,repr(a),repr(b))
    worst[name]=mx
print('formula cells',cnt)
for k,v in worst.items(): print(k,v)
