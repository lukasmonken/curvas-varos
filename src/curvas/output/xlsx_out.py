"""XLSX do dia (Parte 1, seção 14): o layout do Dashboard original, só com valores.

Abas:

- ``Dashboard``: as posições de ``Dashboard!A1:Q31`` da planilha antiga
  (``docs/F0_AUDITORIA.md``), com os valores do LEGADO do dia. O bloco da ETTJ
  (``C37:F57``), que nenhuma fórmula lia, fica de fora;
- ``Corrigido``: tabela anual do CORRIGIDO, CDS do ano corrente, realizado e a
  diferença CORRIGIDO − LEGADO em bps;
- ``Fontes``: datas e situação de cada fonte, os alertas e a identificação da execução.

Nenhuma fórmula: todo número sai pronto do JSON. O openpyxl grava 16 algarismos
significativos (o CSV tem a precisão total). Com ``show_cds=False`` (Q17), nenhuma célula
traz valor de CDS nem a taxa de desconto; os links de coleta (``J9:J17``) ficam. Mesma
entrada, mesmos bytes: as datas do arquivo vêm de ``metadata.as_of_date``.
"""

import io
import zipfile
from collections.abc import Sequence
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import TYPE_CHECKING, Any

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.writer.excel import ExcelWriter

from curvas.calendar import t0_eve
from curvas.config import DEFAULT
from curvas.engine.legacy import MONTH_LABELS, N_YEARS, QUARTER_LABELS, RATE_SHEET_VERTEX_DAYS
from curvas.models import Alert, MonthlyValue, RunOutput
from curvas.output.json_out import hide_cds_change

if TYPE_CHECKING:
    from _typeshed import SizedBuffer, StrPath

Sheet = Any  # abas e células do openpyxl (sem tipos)
Line = Sequence[object]

# Rótulos dos vértices do CDS em Dashboard!H9:H17, na ordem de DEFAULT.urls.investing_cds.
CDS_LABELS: tuple[str, ...] = (
    "6 Months",
    "1 Year",
    "2 Years",
    "3 Years",
    "4 Years",
    "5 Years",
    "7 Years",
    "10 Years",
    "20 Years",
)
ANNUAL_HEADER: tuple[str, ...] = ("Ano", "DI", "Inflação", "CDS", "Juro real", "Taxa de desconto")

PCT = "0.00%"
PCT_FINE = "0.0000%"
BPS_FMT = "0.00"
CURVE_FMT = "0.0000"  # curvas em % como digitadas (Dashboard!C11, F11)
DAYS_FMT = "0"
DATE_FMT = "dd/mm/yyyy"

BOLD = Font(bold=True)
LINK_FONT = Font(color="0563C1", underline="single")
INPUT_FILL = PatternFill("solid", fgColor="D8F0D8")  # "Em verde -> Input" (Dashboard!B7:C7)
YTG_FILL = PatternFill("solid", fgColor="F2F2F2")  # Dashboard!Q3:Q6

# Larguras de Dashboard!A:K da planilha; E foi alargada para "Inflação Implícita".
DASHBOARD_WIDTHS = {
    "A": 9.14,
    "B": 15.43,
    "C": 16.14,
    "D": 2.57,
    "E": 17.0,
    "F": 11.0,
    "G": 7.43,
    "H": 11.0,
    "J": 7.14,
    "K": 11.0,
}


def _text(value: str) -> str:
    """Texto sem caracteres ilegais no XML e que o openpyxl não grava como fórmula."""
    clean = str(ILLEGAL_CHARACTERS_RE.sub("", value))
    return " " + clean if clean.startswith("=") else clean


def _w(
    cell: Sheet,
    value: object = None,
    *,
    fmt: str | None = None,
    bold: bool = False,
    fill: Any = None,
) -> None:
    """Grava um valor (``None`` deixa a célula vazia) e o formato."""
    if value is not None:
        cell.value = _text(value) if isinstance(value, str) else value
    if fmt is not None and not isinstance(value, str):
        cell.number_format = fmt
    if bold:
        cell.font = BOLD
    if fill is not None:
        cell.fill = fill


def _month(month: int | None) -> str | None:
    return MONTH_LABELS[month - 1] if month else None


def _fraction(value: MonthlyValue | None) -> float | None:
    """% no mês → decimal, como no Dashboard original (0,0033 = 0,33%).

    O deslocamento é decimal: 0,7% vira 0,007, e não 0,006999999999999999 de ``0.7 / 100``.
    """
    return None if value is None else float(Decimal(repr(value.value_pct)).scaleb(-2))


def _if(show: bool, value: float | None) -> float | None:
    return value if show else None


# --- Dashboard (LEGADO) ---------------------------------------------------------------


def _dash_params(ws: Sheet, out: RunOutput) -> None:
    """B2:C7. "Mês último DI" = último mês da Selic usada pelo LEGADO."""
    li, selic = out.inputs.legacy, out.realized.selic_monthly_legacy
    _w(ws["B2"], "Trimestre Atual:")
    _w(ws["C2"], QUARTER_LABELS[li.quarter - 1] if li.quarter else None)
    _w(ws["B3"], "Mês último IPCA")
    _w(ws["C3"], _month(li.month), fill=INPUT_FILL)
    _w(ws["B4"], "Mês último DI")
    _w(ws["C4"], _month(selic[-1].month if selic else None), fill=INPUT_FILL)
    _w(ws["B5"], "Atualizado em:")
    _w(ws["C5"], out.metadata.as_of_date, fmt=DATE_FMT)
    _w(ws["B7"], "Em verde ->", bold=True, fill=INPUT_FILL)
    _w(ws["C7"], "Input", bold=True, fill=INPUT_FILL)


def _dash_projections(ws: Sheet, out: RunOutput, show_cds: bool) -> None:
    """E2:Q6: anos em F:P e o YTG em Q. Ano indisponível (None) fica vazio."""
    annual = out.annual
    if len(annual.legacy) != N_YEARS:
        raise ValueError(f"o Dashboard tem {N_YEARS} anos (F:P); a saída tem {len(annual.legacy)}")
    _w(ws["F2"], "Projeções", bold=True)
    ws.merge_cells("F2:Q2")
    for row, label in ((4, "DI"), (5, "CDS"), (6, "Inflação Implícita")):
        _w(ws[f"E{row}"], label)
    for i, year in enumerate(annual.legacy):
        col = get_column_letter(6 + i)
        _w(ws[f"{col}3"], year.year, bold=True)
        _w(ws[f"{col}4"], year.di, fmt=PCT)
        _w(ws[f"{col}5"], _if(show_cds, year.cds), fmt=PCT)
        _w(ws[f"{col}6"], year.inflation, fmt=PCT)
    ytg = annual.legacy_ytg
    _w(ws["Q3"], f"YTG {out.inputs.legacy.first_year}", bold=True, fill=YTG_FILL)
    _w(ws["Q4"], ytg.di, fmt=PCT, fill=YTG_FILL)
    _w(ws["Q5"], _if(show_cds, ytg.cds), fmt=PCT, fill=YTG_FILL)
    _w(ws["Q6"], ytg.inflation, fmt=PCT, fill=YTG_FILL)


def _dash_curves(ws: Sheet, out: RunOutput) -> None:
    """B9:F29: os 19 vértices que o LEGADO lê (B30:C31 não entram em nenhuma fórmula)."""
    li = out.inputs.legacy
    _w(ws["B9"], "Curva de Inflação", bold=True)
    ws.merge_cells("B9:C9")
    _w(ws["E9"], "Curva de Juros", bold=True)
    ws.merge_cells("E9:F9")
    for ref, label in (
        ("B10", "Vértices"),
        ("C10", "Inflação Implícita"),
        ("E10", "Vértices"),
        ("F10", "DI Implícito"),
    ):
        _w(ws[ref], label)
    curves = zip(RATE_SHEET_VERTEX_DAYS, li.inflation_curve_pct, li.di_curve_pct, strict=True)
    for row, (days, inflation, di) in enumerate(curves, start=11):
        _w(ws[f"B{row}"], days, fmt=DAYS_FMT)
        _w(ws[f"C{row}"], inflation, fmt=CURVE_FMT, fill=INPUT_FILL)
        _w(ws[f"E{row}"], days, fmt=DAYS_FMT)
        _w(ws[f"F{row}"], di, fmt=CURVE_FMT, fill=INPUT_FILL)


def _dash_cds(ws: Sheet, out: RunOutput, show_cds: bool) -> None:
    """H8:Q17: bps e o link da aba do Investing.com de cada vértice, para a coleta manual."""
    _w(ws["H8"], "CDS", bold=True)
    _w(ws["I8"], "Valor", bold=True)
    _w(ws["J8"], "Link", bold=True)
    ws.merge_cells("J8:Q8")
    cds = zip(CDS_LABELS, DEFAULT.urls.investing_cds, out.inputs.legacy.cds_bps, strict=True)
    for row, (label, (_, url), bps) in enumerate(cds, start=9):
        _w(ws[f"H{row}"], label)
        _w(ws[f"I{row}"], _if(show_cds, bps), fmt=BPS_FMT, fill=INPUT_FILL)
        link = ws[f"J{row}"]
        _w(link, url)
        link.hyperlink = url
        link.font = LINK_FONT
        ws.merge_cells(f"J{row}:Q{row}")


def _dash_realized(ws: Sheet, out: RunOutput) -> None:
    """H19:L31: IPCA e Selic dos meses 1..m do ano corrente, sem ano, como na planilha.

    O IPCA é o mesmo SGS 433 do realizado do CORRIGIDO, cortado no mês m do LEGADO.
    """
    li = out.inputs.legacy
    last = li.month or 0
    ipca = {
        v.month: v for v in out.realized.ipca_monthly if v.year == li.first_year and v.month <= last
    }
    selic = {v.month: v for v in out.realized.selic_monthly_legacy if v.year == li.first_year}
    _w(ws["H19"], "Inflação Passada", bold=True)
    ws.merge_cells("H19:I19")
    _w(ws["K19"], "Selic Passado", bold=True)
    ws.merge_cells("K19:L19")
    for month, label in enumerate(MONTH_LABELS, start=1):
        row = 19 + month
        _w(ws[f"H{row}"], label)
        _w(ws[f"I{row}"], _fraction(ipca.get(month)), fmt=PCT, fill=INPUT_FILL)
        _w(ws[f"K{row}"], label)
        _w(ws[f"L{row}"], _fraction(selic.get(month)), fmt=PCT, fill=INPUT_FILL)


def _dashboard(ws: Sheet, out: RunOutput, show_cds: bool) -> None:
    _dash_params(ws, out)
    _dash_projections(ws, out, show_cds)
    _dash_curves(ws, out)
    _dash_cds(ws, out, show_cds)
    _dash_realized(ws, out)
    for col, width in DASHBOARD_WIDTHS.items():
        ws.column_dimensions[col].width = width


# --- Abas sequenciais (Corrigido e Fontes) ----------------------------------------------


class _Cursor:
    """Escreve uma aba de cima para baixo, uma linha por chamada, a partir da coluna A."""

    def __init__(self, ws: Sheet) -> None:
        self.ws = ws
        self.row = 0

    def skip(self) -> None:
        self.row += 1

    def title(self, text: str) -> None:
        self.row += 1
        _w(self.ws.cell(row=self.row, column=1), text, bold=True)

    def header(self, labels: Sequence[str]) -> None:
        self.row += 1
        for col, label in enumerate(labels, start=1):
            _w(self.ws.cell(row=self.row, column=col), label, bold=True)

    def line(self, values: Line, fmts: Sequence[str | None] = ()) -> None:
        self.row += 1
        for col, value in enumerate(values, start=1):
            fmt = fmts[col - 1] if col <= len(fmts) else None
            _w(self.ws.cell(row=self.row, column=col), value, fmt=fmt)

    def pairs(self, items: Sequence[tuple[str, object, str | None]]) -> None:
        for label, value, fmt in items:
            self.line([label, value], (None, fmt))


def _corrected_annual(c: _Cursor, out: RunOutput, show_cds: bool) -> None:
    c.title("Premissas anuais: modo CORRIGIDO")
    c.pairs([("Data-base (t0)", out.metadata.as_of_date, DATE_FMT)])
    c.skip()
    c.header(ANNUAL_HEADER)
    for r in out.annual.corrected:
        values = [r.year, r.di, r.inflation, _if(show_cds, r.cds), r.real]
        c.line([*values, _if(show_cds, r.discount)], (None, *[PCT] * 5))
    now = out.annual.corrected_cds_current_year
    c.skip()
    c.title("CDS do ano corrente (E8.7)")
    c.pairs(
        [
            ("Taxa anualizada do período restante", _if(show_cds, now.annualized_rate), PCT),
            (
                "Acumulado no período restante",
                _if(show_cds, now.remaining_period_accumulated),
                PCT_FINE,
            ),
            ("Dias úteis restantes", now.business_days, DAYS_FMT),
        ]
    )


def _corrected_realized(c: _Cursor, out: RunOutput) -> None:
    realized = out.realized
    c.skip()
    c.title("IPCA mensal (realizado)")
    c.header(("Mês", "Ano", "IPCA no mês", "Divulgado em"))
    for v in realized.ipca_monthly:
        c.line([_month(v.month), v.year, _fraction(v), v.release_date], (None, None, PCT, DATE_FMT))
    if not realized.ipca_monthly:
        c.line(["Nenhum IPCA do ano divulgado até t0."])
    c.pairs(
        [
            ("IPCA acumulado no ano", realized.ipca_ytd, PCT),
            ("Regra do IPCA não divulgado", realized.ipca_rule, None),
        ]
    )
    used = realized.ipca15_used
    if used is not None:
        values = ["IPCA-15 usado", f"{_month(used.month)} de {used.year}", _fraction(used)]
        c.line(values, (None, None, PCT))

    c.skip()
    c.title("CDI mensal (realizado)")
    # O mês de t0 entra só até a véspera (Q11-B): marcado como no site (página Realizado).
    t0 = out.metadata.as_of_date
    partial = f"parcial, até {t0_eve(t0):%d/%m/%Y} (véspera de t0)"
    is_partial = [(v.year, v.month) == (t0.year, t0.month) for v in realized.cdi_monthly]
    c.header(("Mês", "Ano", "CDI no mês", *(("Observação",) if any(is_partial) else ())))
    for v, part in zip(realized.cdi_monthly, is_partial, strict=True):
        note = partial if part else None
        c.line([_month(v.month), v.year, _fraction(v), note], (None, None, PCT))
    if not realized.cdi_monthly:
        c.line(["Sem CDI mensal nesta saída."])
    cdi = realized.cdi
    c.pairs(
        [
            ("CDI acumulado no ano", cdi.factor - 1, PCT),
            ("CDI: início", cdi.start, DATE_FMT),
            ("CDI: até a véspera de", cdi.end_exclusive, DATE_FMT),
            ("CDI: dias úteis", cdi.business_days, DAYS_FMT),
            ("CDI: última observação antes de t0", cdi.last_observation, DATE_FMT),
        ]
    )


def _corrected_comparison(c: _Cursor, out: RunOutput, show_cds: bool) -> None:
    c.skip()
    c.title("CORRIGIDO − LEGADO (bps)")
    c.header(ANNUAL_HEADER)
    for r in out.comparison:
        values = [r.year, r.di_bps, r.inflation_bps, _if(show_cds, r.cds_bps), r.real_bps]
        c.line([*values, _if(show_cds, r.discount_bps)], (None, *[BPS_FMT] * 5))


def _corrected(ws: Sheet, out: RunOutput, show_cds: bool) -> None:
    c = _Cursor(ws)
    _corrected_annual(c, out, show_cds)
    _corrected_realized(c, out)
    _corrected_comparison(c, out, show_cds)
    ws.column_dimensions["A"].width = 36
    for col in "BCDEF":
        ws.column_dimensions[col].width = 16


def _alerts(out: RunOutput, show_cds: bool) -> list[Alert]:
    """Com ``show_cds=False``, a variação diária do CDS e da taxa de desconto sai sem o valor."""
    if show_cds:
        return list(out.alerts)
    return [hide_cds_change(a) for a in out.alerts]


def _when(value: datetime | date | None) -> object:
    # O Excel não guarda fuso: data e hora com fuso vão como texto ISO.
    return value.isoformat() if isinstance(value, datetime) else value


def _sources(ws: Sheet, out: RunOutput, show_cds: bool) -> None:
    c = _Cursor(ws)
    c.header(
        (
            "Fonte",
            "Data pedida",
            "Data do dado",
            "Defasada",
            "Motivo do fallback",
            "Publicado em",
            "Baixado em",
            "URL",
        )
    )
    for s in out.sources:
        values = [
            s.source,
            s.requested_date,
            s.source_date,
            "sim" if s.stale else "não",
            s.fallback_reason,
            _when(s.publication_date),
            _when(s.retrieved_at),
            s.url,
        ]
        c.line(values, (None, DATE_FMT, DATE_FMT, None, None, DATE_FMT))

    c.skip()
    c.title("Alertas")
    c.header(("Nível", "Código", "Mensagem", "Fonte"))
    alerts = _alerts(out, show_cds)
    for a in alerts:
        c.line([a.level, a.code, a.message, a.source])
    if not alerts:
        c.line(["Nenhum alerta."])

    meta = out.metadata
    c.skip()
    c.title("Execução")
    c.pairs(
        [
            ("Data-base (t0)", meta.as_of_date, DATE_FMT),
            ("Versão do código", meta.code_version, None),
            ("Commit", meta.git_commit, None),
            ("Versão do schema", meta.schema_version, None),
        ]
    )
    for col, width in zip("ABCDEFGH", (18, 12, 12, 10, 60, 26, 26, 60), strict=True):
        ws.column_dimensions[col].width = width


# --- Gravação determinística ------------------------------------------------------------


class _FixedDateZip(zipfile.ZipFile):
    """ZipFile que grava toda entrada com a mesma data, para bytes determinísticos.

    O openpyxl grava as partes com ``writestr`` e as abas com ``write`` (de um arquivo
    temporário); nos dois casos o zip levaria a hora atual.
    """

    def __init__(self, file: io.BytesIO, date_time: tuple[int, int, int, int, int, int]) -> None:
        super().__init__(file, "w", zipfile.ZIP_DEFLATED)
        self._date_time = date_time

    def writestr(
        self,
        zinfo_or_arcname: "str | zipfile.ZipInfo",
        data: "SizedBuffer | str",
        compress_type: int | None = None,
        compresslevel: int | None = None,
    ) -> None:
        name = (
            zinfo_or_arcname.filename
            if isinstance(zinfo_or_arcname, zipfile.ZipInfo)
            else zinfo_or_arcname
        )
        info = zipfile.ZipInfo(name, date_time=self._date_time)
        info.compress_type = zipfile.ZIP_DEFLATED
        info.external_attr = 0o600 << 16
        super().writestr(info, data)

    def write(
        self,
        filename: "StrPath",
        arcname: "StrPath | None" = None,
        compress_type: int | None = None,
        compresslevel: int | None = None,
    ) -> None:
        path = Path(filename)
        self.writestr(path.name if arcname is None else str(arcname), path.read_bytes())


def build_day_xlsx(out: RunOutput, *, show_cds: bool = True) -> bytes:
    """XLSX do dia: abas Dashboard (LEGADO), Corrigido e Fontes, só com valores."""
    wb = Workbook()
    dashboard = wb.active
    dashboard.title = "Dashboard"
    _dashboard(dashboard, out, show_cds)
    _corrected(wb.create_sheet("Corrigido"), out, show_cds)
    _sources(wb.create_sheet("Fontes"), out, show_cds)

    as_of = out.metadata.as_of_date
    stamp = datetime(as_of.year, as_of.month, as_of.day)
    wb.properties.creator = "curvas"
    wb.properties.title = f"Curvas VAROS {as_of:%d/%m/%Y}"
    wb.properties.created = stamp
    wb.properties.modified = stamp  # o Workbook.save() poria a hora atual aqui

    buf = io.BytesIO()
    ExcelWriter(wb, _FixedDateZip(buf, (as_of.year, as_of.month, as_of.day, 0, 0, 0))).save()
    return buf.getvalue()
