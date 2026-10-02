"""Downloads do site (Parte 1, seção 14): CSV e XLSX do dia, só com valores."""

import csv
import io
import math
import re
import zipfile
from datetime import date, datetime, timedelta
from typing import Any

import openpyxl
import pytest

from conftest import ROOT, XLSX
from curvas.config import DEFAULT
from curvas.engine.legacy import MONTH_LABELS, QUARTER_LABELS, RATE_SHEET_VERTEX_DAYS
from curvas.models import Alert, ComparisonRow, LegacyYtg, MonthlyValue, RunOutput
from curvas.output.csv_out import render_day_csv, render_history_csv
from curvas.output.xlsx_out import CDS_LABELS, build_day_xlsx

# Cópia congelada (o data/latest.json de produção muda todo dia; ver tests/test_site.py).
LATEST = ROOT / "tests" / "fixtures" / "f5" / "data" / "latest.json"
FIELDS = (("di", "di"), ("inflacao", "inflation"), ("juro_real", "real"))
CDS_FIELDS = (("cds", "cds"), ("desconto", "discount"))
# Hosts que podem aparecer no XLSX: fontes do pipeline, abas do CDS e namespaces do
# Office Open XML. Qualquer outro (por exemplo, um link interno copiado da planilha) falha.
ALLOWED_HOSTS = {
    "api.bcb.gov.br",
    "br.investing.com",
    "purl.org",
    "schemas.openxmlformats.org",
    "servicodados.ibge.gov.br",
    "www.anbima.com.br",
    "www.w3.org",
}


@pytest.fixture(scope="module")
def latest() -> RunOutput:
    return RunOutput.model_validate_json(LATEST.read_text(encoding="utf-8"))


def _revalidate(out: RunOutput) -> RunOutput:
    return RunOutput.model_validate(out.model_dump())


def _with_date(out: RunOutput, day: date) -> RunOutput:
    meta = out.metadata.model_copy(update={"as_of_date": day})
    return _revalidate(out.model_copy(update={"metadata": meta}))


@pytest.fixture(scope="module")
def virada(latest: RunOutput) -> RunOutput:
    """Saída sintética da virada do ano (Q6): ano corrente do LEGADO indisponível.

    Também sem CDI mensal e com alertas de variação diária (um deles do CDS) e um
    texto que começa com "=".
    """
    empty = {"di": None, "inflation": None, "cds": None, "real": None, "discount": None}
    legacy = [latest.annual.legacy[0].model_copy(update=empty), *latest.annual.legacy[1:]]
    annual = latest.annual.model_copy(
        update={"legacy": legacy, "legacy_ytg": LegacyYtg(di=None, inflation=None, cds=None)}
    )
    li = latest.inputs.legacy.model_copy(update={"month": None, "quarter": None})
    first = latest.comparison[0].year
    comparison = [
        ComparisonRow(
            year=first,
            di_bps=None,
            inflation_bps=None,
            cds_bps=None,
            real_bps=None,
            discount_bps=None,
        ),
        *latest.comparison[1:],
    ]
    alerts = [
        *latest.alerts,
        Alert(
            level="warning",
            code="variacao_diaria",
            message="corrected cds 2027: +60.0 bps desde 2026-09-30 (limite 50 bps)",
        ),
        Alert(
            level="warning",
            code="variacao_diaria",
            message="legacy di 2028: -55.0 bps desde 2026-09-30 (limite 50 bps)",
        ),
        Alert(level="info", code="teste", message="=1+1 não é fórmula"),
    ]
    out = latest.model_copy(
        update={
            "annual": annual,
            "inputs": latest.inputs.model_copy(update={"legacy": li}),
            "realized": latest.realized.model_copy(
                update={"selic_monthly_legacy": [], "cdi_monthly": []}
            ),
            "comparison": comparison,
            "alerts": alerts,
        }
    )
    return _revalidate(out)


def _close(got: Any, expected: float | None) -> bool:
    """O openpyxl grava 16 algarismos significativos; o CSV, a precisão total."""
    if expected is None:
        return got is None
    return isinstance(got, int | float) and math.isclose(got, expected, rel_tol=1e-15)


def _load(data: bytes) -> Any:
    return openpyxl.load_workbook(io.BytesIO(data))


def _cells(wb: Any) -> list[Any]:
    return [c for ws in wb for row in ws.iter_rows() for c in row]


def _find(ws: Any, text: str) -> int:
    for (cell,) in ws.iter_rows(min_col=1, max_col=1):
        if cell.value == text:
            return int(cell.row)
    raise AssertionError(f"{text!r} não encontrado na aba {ws.title}")


def _cds_values(out: RunOutput) -> list[float]:
    """Todo número que é CDS ou depende dele (taxa de desconto)."""
    a = out.annual
    values: list[float | None] = [
        *(r.cds for r in [*a.corrected, *a.legacy]),
        *(r.discount for r in [*a.corrected, *a.legacy]),
        a.legacy_ytg.cds,
        a.corrected_cds_current_year.annualized_rate,
        a.corrected_cds_current_year.remaining_period_accumulated,
        *out.inputs.legacy.cds_bps,
        *(c.cds_bps for c in out.comparison),
        *(c.discount_bps for c in out.comparison),
    ]
    return [v for v in values if v is not None]


# --- CSV ----------------------------------------------------------------------------------


def _check_csv_rows(rows: list[dict[str, str]], out: RunOutput, show_cds: bool) -> None:
    fields = FIELDS + (CDS_FIELDS if show_cds else ())
    expected = [("corrigido", r) for r in out.annual.corrected] + [
        ("legado", r) for r in out.annual.legacy
    ]
    assert len(rows) == len(expected)
    for got, (mode, row) in zip(rows, expected, strict=True):
        assert got["data_base"] == out.metadata.as_of_date.isoformat()
        assert got["modo"] == mode
        assert int(got["ano"]) == row.year
        for column, field in fields:
            value = getattr(row, field)
            # repr() ida e volta: igualdade exata, sem tolerância.
            assert got[column] == ("" if value is None else repr(value))
            assert (None if got[column] == "" else float(got[column])) == value


def test_csv_dia_bate_com_json(latest: RunOutput) -> None:
    text = render_day_csv(latest)
    assert "\r" not in text
    assert text.endswith("\n")
    reader = csv.DictReader(io.StringIO(text))
    assert reader.fieldnames == [
        "data_base",
        "modo",
        "ano",
        "di",
        "inflacao",
        "cds",
        "juro_real",
        "desconto",
    ]
    _check_csv_rows(list(reader), latest, show_cds=True)


def test_csv_sem_cds(latest: RunOutput) -> None:
    reader = csv.DictReader(io.StringIO(render_day_csv(latest, show_cds=False)))
    assert reader.fieldnames == ["data_base", "modo", "ano", "di", "inflacao", "juro_real"]
    _check_csv_rows(list(reader), latest, show_cds=False)


def test_csv_valor_indisponivel_fica_vazio(virada: RunOutput) -> None:
    rows = list(csv.DictReader(io.StringIO(render_day_csv(virada))))
    first = virada.annual.legacy[0].year
    legacy_first = next(r for r in rows if r["modo"] == "legado" and int(r["ano"]) == first)
    assert all(legacy_first[c] == "" for c in ("di", "inflacao", "cds", "juro_real", "desconto"))
    _check_csv_rows(rows, virada, show_cds=True)


def test_csv_historico(latest: RunOutput) -> None:
    day = latest.metadata.as_of_date
    outputs = [_with_date(latest, day - timedelta(days=k)) for k in (0, 2, 1)]
    text = render_history_csv(outputs)
    assert text.splitlines()[0] == "data_base,modo,ano,di,inflacao,cds,juro_real,desconto"
    rows = list(csv.DictReader(io.StringIO(text)))
    per_day = 2 * len(latest.annual.corrected)
    assert len(rows) == 3 * per_day
    ordered = sorted(outputs, key=lambda o: o.metadata.as_of_date)
    for i, out in enumerate(ordered):
        _check_csv_rows(rows[i * per_day : (i + 1) * per_day], out, show_cds=True)
    keys = [(r["data_base"], r["modo"], r["ano"]) for r in rows]
    assert len(set(keys)) == len(keys)
    # O histórico de uma data só é o próprio CSV do dia.
    assert render_history_csv([latest], show_cds=False) == render_day_csv(latest, show_cds=False)


def test_csv_historico_recusa_data_repetida(latest: RunOutput) -> None:
    with pytest.raises(ValueError, match="repetida"):
        render_history_csv([latest, latest])


# --- XLSX: estrutura, fórmulas, link interno e determinismo ------------------------------


@pytest.fixture(scope="module", params=["latest", "virada"])
def any_output(request: pytest.FixtureRequest, latest: RunOutput, virada: RunOutput) -> RunOutput:
    return latest if request.param == "latest" else virada


@pytest.mark.parametrize("show_cds", [True, False])
def test_xlsx_sem_formulas(any_output: RunOutput, show_cds: bool) -> None:
    data = build_day_xlsx(any_output, show_cds=show_cds)
    wb = _load(data)
    assert wb.sheetnames == ["Dashboard", "Corrigido", "Fontes"]
    for cell in _cells(wb):
        assert cell.data_type != "f", cell.coordinate
        assert not (isinstance(cell.value, str) and cell.value.startswith("=")), cell.coordinate
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        for name in z.namelist():
            if name.startswith("xl/worksheets/sheet"):
                xml = z.read(name)
                assert b"<f>" not in xml
                assert b"<f " not in xml


@pytest.mark.parametrize("show_cds", [True, False])
def test_xlsx_sem_link_interno(any_output: RunOutput, show_cds: bool) -> None:
    data = build_day_xlsx(any_output, show_cds=show_cds)
    wb = _load(data)
    targets = set()
    for cell in _cells(wb):
        if cell.hyperlink is not None:
            targets.add(cell.hyperlink.target)
    # Os únicos hyperlinks são as abas do Investing.com da coleta manual (Q17).
    assert targets == {url for _, url in DEFAULT.urls.investing_cds}
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        for name in z.namelist():
            hosts = {h.decode() for h in re.findall(rb"https?://([A-Za-z0-9.-]+)", z.read(name))}
            assert hosts <= ALLOWED_HOSTS, (name, hosts - ALLOWED_HOSTS)


@pytest.mark.parametrize("show_cds", [True, False])
def test_xlsx_deterministico(any_output: RunOutput, show_cds: bool) -> None:
    first = build_day_xlsx(any_output, show_cds=show_cds)
    assert build_day_xlsx(any_output, show_cds=show_cds) == first
    day = any_output.metadata.as_of_date
    with zipfile.ZipFile(io.BytesIO(first)) as z:
        assert {i.date_time for i in z.infolist()} == {(day.year, day.month, day.day, 0, 0, 0)}
    props = _load(first).properties
    stamp = datetime(day.year, day.month, day.day)
    assert props.created == stamp
    assert props.modified == stamp


# --- XLSX: Dashboard (LEGADO) --------------------------------------------------------------


def test_xlsx_dashboard_valores(latest: RunOutput) -> None:
    ws = _load(build_day_xlsx(latest))["Dashboard"]
    li, annual = latest.inputs.legacy, latest.annual
    assert ws["B2"].value == "Trimestre Atual:"
    # data/latest.json muda todo dia: na virada do ano, mês e trimestre ficam vazios (Q6).
    selic_months = [v.month for v in latest.realized.selic_monthly_legacy]
    assert ws["C2"].value == (QUARTER_LABELS[li.quarter - 1] if li.quarter else None)
    assert ws["C3"].value == (MONTH_LABELS[li.month - 1] if li.month else None)
    assert ws["C4"].value == (MONTH_LABELS[selic_months[-1] - 1] if selic_months else None)
    assert ws["C5"].value.date() == latest.metadata.as_of_date
    assert ws["F2"].value == "Projeções"
    assert ws["Q3"].value == f"YTG {li.first_year}"
    assert (ws["E4"].value, ws["E5"].value, ws["E6"].value) == ("DI", "CDS", "Inflação Implícita")
    for i, row in enumerate(annual.legacy):
        col = openpyxl.utils.get_column_letter(6 + i)
        assert ws[f"{col}3"].value == row.year
        assert _close(ws[f"{col}4"].value, row.di)
        assert _close(ws[f"{col}5"].value, row.cds)
        assert _close(ws[f"{col}6"].value, row.inflation)
        assert ws[f"{col}4"].number_format == "0.00%"
    ytg = annual.legacy_ytg
    assert _close(ws["Q4"].value, ytg.di)
    assert _close(ws["Q5"].value, ytg.cds)
    assert _close(ws["Q6"].value, ytg.inflation)
    for k, days in enumerate(RATE_SHEET_VERTEX_DAYS):
        assert ws[f"B{11 + k}"].value == days
        assert ws[f"E{11 + k}"].value == days
        assert _close(ws[f"C{11 + k}"].value, li.inflation_curve_pct[k])
        assert _close(ws[f"F{11 + k}"].value, li.di_curve_pct[k])
    for k, bps in enumerate(li.cds_bps):
        assert ws[f"H{9 + k}"].value == CDS_LABELS[k]
        assert _close(ws[f"I{9 + k}"].value, bps)
    last = li.month or 0
    ipca = {v.month: v.value_pct for v in latest.realized.ipca_monthly if v.month <= last}
    selic = {v.month: v.value_pct for v in latest.realized.selic_monthly_legacy}
    for m, label in enumerate(MONTH_LABELS, start=1):
        assert ws[f"H{19 + m}"].value == label
        assert ws[f"K{19 + m}"].value == label
        expected_ipca = ipca[m] / 100 if m in ipca else None
        expected_selic = selic[m] / 100 if m in selic else None
        assert _close(ws[f"I{19 + m}"].value, expected_ipca)
        assert _close(ws[f"L{19 + m}"].value, expected_selic)
        assert ws[f"I{19 + m}"].number_format == "0.00%"


@pytest.mark.parametrize("show_cds", [True, False])
def test_xlsx_links_do_cds(latest: RunOutput, show_cds: bool) -> None:
    """Q17: os links de coleta manual (J9:J17) ficam mesmo sem os valores."""
    ws = _load(build_day_xlsx(latest, show_cds=show_cds))["Dashboard"]
    assert (ws["H8"].value, ws["I8"].value, ws["J8"].value) == ("CDS", "Valor", "Link")
    urls = [url for _, url in DEFAULT.urls.investing_cds]
    assert len(urls) == 9
    for k, url in enumerate(urls):
        cell = ws[f"J{9 + k}"]
        assert cell.value == url
        assert cell.hyperlink is not None
        assert cell.hyperlink.target == url
        assert url.startswith("https://br.investing.com/")


def test_xlsx_virada_do_ano(virada: RunOutput) -> None:
    """Ano corrente do LEGADO indisponível (Q6): células vazias, anos seguintes normais."""
    wb = _load(build_day_xlsx(virada))
    ws = wb["Dashboard"]
    for ref in ("C2", "C3", "C4", "F4", "F5", "F6", "Q4", "Q5", "Q6"):
        assert ws[ref].value is None, ref
    assert _close(ws["G4"].value, virada.annual.legacy[1].di)
    assert all(ws[f"I{r}"].value is None for r in range(20, 32))
    assert all(ws[f"L{r}"].value is None for r in range(20, 32))
    corr = wb["Corrigido"]
    row = _find(corr, "CORRIGIDO − LEGADO (bps)") + 2
    assert corr.cell(row=row, column=1).value == virada.comparison[0].year
    assert all(corr.cell(row=row, column=c).value is None for c in range(2, 7))
    assert corr.cell(row=_find(corr, "CDI mensal (realizado)") + 2, column=1).value == (
        "Sem CDI mensal nesta saída."
    )


# --- XLSX: sem CDS (Q17) -------------------------------------------------------------------


def test_xlsx_sem_cds_nenhum_valor(latest: RunOutput) -> None:
    wb = _load(build_day_xlsx(latest, show_cds=False))
    forbidden = _cds_values(latest)
    assert forbidden
    for cell in _cells(wb):
        # Inteiros são anos e contagens de dias (126, 252…, 62, 187): um CDS de exatamente
        # 126,00 bps colidiria sem vazar nada. As células que levariam CDS, inteiro ou não,
        # são conferidas uma a uma abaixo.
        if isinstance(cell.value, float):
            hit = [v for v in forbidden if math.isclose(cell.value, v, rel_tol=1e-15)]
            assert not hit, f"{cell.parent.title}!{cell.coordinate} = {cell.value}"
    ws = wb["Dashboard"]
    for ref in [f"{openpyxl.utils.get_column_letter(c)}5" for c in range(6, 18)]:
        assert ws[ref].value is None, ref
    assert all(ws[f"I{r}"].value is None for r in range(9, 18))
    assert ws["E5"].value == "CDS"  # rótulos ficam
    corr = wb["Corrigido"]
    top = _find(corr, "Ano")
    for r in range(top + 1, top + 1 + len(latest.annual.corrected)):
        assert corr.cell(row=r, column=4).value is None
        assert corr.cell(row=r, column=6).value is None
        assert corr.cell(row=r, column=2).value is not None
    for label in ("Taxa anualizada do período restante", "Acumulado no período restante"):
        assert corr.cell(row=_find(corr, label), column=2).value is None
    top = _find(corr, "CORRIGIDO − LEGADO (bps)") + 1
    for r in range(top + 1, top + 1 + len(latest.comparison)):
        assert corr.cell(row=r, column=4).value is None
        assert corr.cell(row=r, column=6).value is None
        assert corr.cell(row=r, column=2).value is not None


def test_xlsx_sem_cds_tira_o_valor_da_variacao_do_cds(virada: RunOutput) -> None:
    def messages(show_cds: bool) -> list[str]:
        ws = _load(build_day_xlsx(virada, show_cds=show_cds))["Fontes"]
        start = _find(ws, "Alertas") + 2
        return [
            str(ws.cell(row=r, column=3).value)
            for r in range(start, ws.max_row + 1)
            if ws.cell(row=r, column=2).value
        ]

    assert any(m.startswith("corrected cds 2027: +60.0") for m in messages(True))
    hidden = messages(False)
    # Mesma regra do site: o alerta fica, sem o valor.
    assert "corrected cds 2027: variação acima do limite" in hidden
    assert not any("+60.0" in m for m in hidden)
    assert any(m.startswith("legacy di 2028: -55.0") for m in hidden)
    # Texto que começa com "=" sai como texto, nunca como fórmula.
    assert " =1+1 não é fórmula" in hidden


# --- XLSX: Corrigido e Fontes --------------------------------------------------------------


def test_xlsx_corrigido(latest: RunOutput) -> None:
    ws = _load(build_day_xlsx(latest))["Corrigido"]
    top = _find(ws, "Ano")
    assert [ws.cell(row=top, column=c).value for c in range(1, 7)] == [
        "Ano",
        "DI",
        "Inflação",
        "CDS",
        "Juro real",
        "Taxa de desconto",
    ]
    for k, r in enumerate(latest.annual.corrected, start=top + 1):
        values = [ws.cell(row=k, column=c).value for c in range(1, 7)]
        assert values[0] == r.year
        expected = (r.di, r.inflation, r.cds, r.real, r.discount)
        assert all(_close(g, e) for g, e in zip(values[1:], expected, strict=True))

    now = latest.annual.corrected_cds_current_year
    row = _find(ws, "CDS do ano corrente (E8.7)")
    assert _close(ws.cell(row=row + 1, column=2).value, now.annualized_rate)
    assert _close(ws.cell(row=row + 2, column=2).value, now.remaining_period_accumulated)
    assert ws.cell(row=row + 3, column=2).value == now.business_days

    row = _find(ws, "IPCA mensal (realizado)") + 1
    for k, v in enumerate(latest.realized.ipca_monthly, start=row + 1):
        assert ws.cell(row=k, column=1).value == MONTH_LABELS[v.month - 1]
        assert ws.cell(row=k, column=2).value == v.year
        assert _close(ws.cell(row=k, column=3).value, v.value_pct / 100)
    ytd = _find(ws, "IPCA acumulado no ano")
    assert _close(ws.cell(row=ytd, column=2).value, latest.realized.ipca_ytd)
    assert ws.cell(row=_find(ws, "Regra do IPCA não divulgado"), column=2).value == (
        latest.realized.ipca_rule
    )
    cdi = _find(ws, "CDI acumulado no ano")
    assert _close(ws.cell(row=cdi, column=2).value, latest.realized.cdi.factor - 1)

    row = _find(ws, "CORRIGIDO − LEGADO (bps)") + 1
    for k, c in enumerate(latest.comparison, start=row + 1):
        assert ws.cell(row=k, column=1).value == c.year
        expected = (c.di_bps, c.inflation_bps, c.cds_bps, c.real_bps, c.discount_bps)
        got = [ws.cell(row=k, column=col).value for col in range(2, 7)]
        assert all(_close(g, e) for g, e in zip(got, expected, strict=True))


def test_xlsx_corrigido_cdi_mensal(latest: RunOutput) -> None:
    months = [MonthlyValue(year=2026, month=m, value_pct=1.0 + m / 100) for m in (1, 2, 3)]
    out = _revalidate(
        latest.model_copy(
            update={"realized": latest.realized.model_copy(update={"cdi_monthly": months})}
        )
    )
    ws = _load(build_day_xlsx(out))["Corrigido"]
    row = _find(ws, "CDI mensal (realizado)") + 1
    # t0 = 01/10/2026: nenhum mês parcial, sem a coluna de observação.
    header = [ws.cell(row=row, column=c).value for c in (1, 2, 3, 4)]
    assert header == ["Mês", "Ano", "CDI no mês", None]
    for k, v in enumerate(months, start=row + 1):
        assert ws.cell(row=k, column=1).value == MONTH_LABELS[v.month - 1]
        assert ws.cell(row=k, column=2).value == v.year
        assert _close(ws.cell(row=k, column=3).value, v.value_pct / 100)
        assert ws.cell(row=k, column=4).value is None
    assert ws.cell(row=row + len(months) + 1, column=1).value == "CDI acumulado no ano"
    last = _find(ws, "CDI: última observação antes de t0")
    assert ws.cell(row=last, column=2).value.date() == latest.realized.cdi.last_observation


def test_xlsx_corrigido_cdi_do_mes_de_t0_parcial(latest: RunOutput) -> None:
    """t0 no meio do mês: o CDI desse mês vai só até a véspera, marcado como no site."""
    months = [MonthlyValue(year=2026, month=m, value_pct=1.0) for m in (8, 9)]
    meta = latest.metadata.model_copy(update={"as_of_date": date(2026, 9, 29)})
    out = _revalidate(
        latest.model_copy(
            update={
                "metadata": meta,
                "realized": latest.realized.model_copy(update={"cdi_monthly": months}),
            }
        )
    )
    ws = _load(build_day_xlsx(out))["Corrigido"]
    row = _find(ws, "CDI mensal (realizado)") + 1
    header = [ws.cell(row=row, column=c).value for c in (1, 2, 3, 4)]
    assert header == ["Mês", "Ano", "CDI no mês", "Observação"]
    assert ws.cell(row=row + 1, column=4).value is None  # agosto: mês inteiro
    assert ws.cell(row=row + 2, column=1).value == "Setembro"
    assert ws.cell(row=row + 2, column=4).value == "parcial, até 28/09/2026 (véspera de t0)"


def test_xlsx_fontes(latest: RunOutput) -> None:
    ws = _load(build_day_xlsx(latest))["Fontes"]
    header = [ws.cell(row=1, column=c).value for c in range(1, 6)]
    assert header == ["Fonte", "Data pedida", "Data do dado", "Defasada", "Motivo do fallback"]
    for k, s in enumerate(latest.sources, start=2):
        assert ws.cell(row=k, column=1).value == s.source
        assert ws.cell(row=k, column=2).value.date() == s.requested_date
        got = ws.cell(row=k, column=3).value
        assert (None if got is None else got.date()) == s.source_date
        assert ws.cell(row=k, column=4).value == ("sim" if s.stale else "não")
        assert ws.cell(row=k, column=5).value == s.fallback_reason
    row = _find(ws, "Alertas") + 1
    if latest.alerts:
        for k, a in enumerate(latest.alerts, start=row + 1):
            assert ws.cell(row=k, column=2).value == a.code
            assert ws.cell(row=k, column=3).value == a.message
    else:
        assert ws.cell(row=row + 1, column=1).value == "Nenhum alerta."


# --- XLSX: posições do Dashboard original (planilha local) ---------------------------------

VARIABLE = {"C2", "C3", "C4", "Q3"}  # dependem da data-base


@pytest.fixture(scope="module")
def original() -> Any:
    if not XLSX.exists():
        pytest.skip("planilha fora do repositório (confidencial)")
    return openpyxl.load_workbook(XLSX)["Dashboard"]


def _strings(ws: Any) -> dict[str, str]:
    return {
        c.coordinate: c.value
        for row in ws.iter_rows(min_row=1, max_row=31, max_col=17)
        for c in row
        if c.data_type == "s" and c.coordinate not in VARIABLE
    }


def test_dashboard_nas_celulas_do_original(latest: RunOutput, original: Any) -> None:
    ws = _load(build_day_xlsx(latest))["Dashboard"]
    assert _strings(ws) == _strings(original)
    filled = {
        c.coordinate
        for row in original.iter_rows(min_row=1, max_row=40, max_col=17)
        for c in row
        if c.value is not None
    }
    ours = {c.coordinate for row in ws.iter_rows() for c in row if c.value is not None}
    assert ours <= filled, sorted(ours - filled)
    assert {str(r) for r in ws.merged_cells.ranges} == {
        str(r) for r in original.merged_cells.ranges
    }
    for k in range(9, 18):
        assert ws[f"J{k}"].hyperlink.target == original[f"J{k}"].hyperlink.target
    # Percentual onde o original mostra taxa.
    for ref in ("F4", "P4", "Q4", "F5", "Q5", "F6", "Q6", "I20", "I31", "L20", "L31"):
        assert ws[ref].number_format == original[ref].number_format == "0.00%", ref
