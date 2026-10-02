"""Gera as fixtures do modo LEGADO a partir da planilha recalculada.

Referência de recálculo: motor ``formulas`` (decisão Q7 em ``OPEN_QUESTIONS.md``).
O script não faz parte do runtime. Uso:

    uv run --group tools python tools/gerar_fixtures_legado.py [--inputs-only]

Saídas em ``tests/fixtures/legacy/``:

- ``inputs_planilha.json``: inputs exatamente como estão na planilha;
- ``esperado_planilha.json``: saídas recalculadas, com as grades diárias completas;
- ``esperado_cenarios.json``: saídas recalculadas com mês/trimestre/YTG alterados.
"""

from __future__ import annotations

import hashlib
import json
import sys
from importlib.metadata import version
from pathlib import Path
from typing import Any

import formulas
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "docs" / "legado" / "Query-Avila.xlsx"
OUT = ROOT / "tests" / "fixtures" / "legacy"
BOOK = "[Query-Avila.xlsx]"

DI_LAST_ROW = 3000  # DI!W3000 = dia 2998
INF_LAST_ROW = 3002  # Inflação!W3002 = dia 3000
CDS_LAST_ROW = 5047  # CDS!L5047 = dia 5040
COLS_E_TO_O = "EFGHIJKLMNO"
COLS_F_TO_P = "FGHIJKLMNOP"

# Cenários além da planilha original. Cada um altera só parâmetros de controle.
SCENARIOS: list[dict[str, Any]] = [
    {
        "id": "agosto_3tri",
        "descricao": "Inflação e DI em Agosto/3º Tri (YTG 84; expoente do DI!E4 = 2)",
        "inflation_params": ["Agosto", "3º Tri"],
        "di_params": ["Agosto", "3º Tri"],
    },
    {
        "id": "setembro_3tri",
        "descricao": "Inflação e DI em Setembro/3º Tri (YTG 63)",
        "inflation_params": ["Setembro", "3º Tri"],
        "di_params": ["Setembro", "3º Tri"],
    },
    {
        "id": "marco_1tri",
        "descricao": "Março/1º Tri: todo o realizado entra em T (dia 1)",
        "inflation_params": ["Março", "1º Tri"],
        "di_params": ["Março", "1º Tri"],
    },
    {
        "id": "marco_2tri",
        "descricao": "Março/2º Tri: T vazio, todo o realizado entra em U; expoente 4/3",
        "inflation_params": ["Março", "2º Tri"],
        "di_params": ["Março", "2º Tri"],
    },
    {
        "id": "novembro_4tri",
        "descricao": "Novembro/4º Tri: YTG 21; expoente do DI!E4 = 4",
        "inflation_params": ["Novembro", "4º Tri"],
        "di_params": ["Novembro", "4º Tri"],
    },
    {
        "id": "abas_diferentes",
        "descricao": "Inflação em Agosto/3º Tri e DI em Setembro/3º Tri (abas independentes)",
        "inflation_params": ["Agosto", "3º Tri"],
        "di_params": ["Setembro", "3º Tri"],
    },
    {
        "id": "inconsistente_maio_3tri",
        "descricao": "Maio/3º Tri (mês anterior ao trimestre): T vazio, U = Jan…Jun",
        "inflation_params": ["Maio", "3º Tri"],
        "di_params": ["Maio", "3º Tri"],
    },
    {
        "id": "cds_ytg_200",
        "descricao": "CDS!I9 = 200 (> 1º vértice): CDS!E4 deixa de ser o 1º vértice",
        "cds_ytg_days": 200,
    },
    {
        "id": "cds_ytg_21",
        "descricao": "CDS!I9 = 21",
        "cds_ytg_days": 21,
    },
    {
        "id": "dezembro_4tri",
        "descricao": "Dezembro/4º Tri: YTG = 0, VLOOKUP(0) dá #N/A só nas células dependentes",
        "inflation_params": ["Dezembro", "4º Tri"],
        "di_params": ["Dezembro", "4º Tri"],
    },
    {
        "id": "dezembro_so_inflacao",
        "descricao": "Só a Inflação em Dezembro: DI e CDS seguem válidos",
        "inflation_params": ["Dezembro", "4º Tri"],
    },
    {
        "id": "cds_ytg_0",
        "descricao": "CDS!I9 = 0: #N/A em J9, J10, E4 e F4; o resto do CDS segue válido",
        "cds_ytg_days": 0,
    },
    {
        "id": "cds_ytg_3000",
        "descricao": "CDS!I9 = 3000: os dois últimos cortes passam do dia 5040 (#N/A)",
        "cds_ytg_days": 3000,
    },
]


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read_inputs() -> dict[str, Any]:
    """Inputs constantes da planilha, lidos sem recálculo."""
    wb = openpyxl.load_workbook(XLSX, data_only=False)
    dash, di, inf, cds = wb["Dashboard"], wb["DI"], wb["Inflação"], wb["CDS"]
    # Prazos dos vértices são fórmulas inteiras (=B15+126, =252*2): o cache é exato.
    cached = openpyxl.load_workbook(XLSX, data_only=True)

    def col(ws: Any, letter: str, first: int, last: int) -> list[float]:
        return [float(ws[f"{letter}{r}"].value) for r in range(first, last + 1)]

    def days(ws: Any, letter: str, first: int, last: int) -> list[int]:
        values = [ws[f"{letter}{r}"].value for r in range(first, last + 1)]
        assert all(float(v) == int(v) for v in values), "prazo não inteiro"
        return [int(v) for v in values]

    return {
        "first_year": int(dash["F3"].value),
        "inflation_curve_pct": col(dash, "C", 11, 29),
        "di_curve_pct": col(dash, "F", 11, 29),
        "cds_bps": col(dash, "I", 9, 17),
        "ipca_monthly": col(dash, "I", 20, 31),
        "selic_monthly": col(dash, "L", 20, 31),
        "inflation_params": [inf["C2"].value, inf["C3"].value],
        "di_params": [di["C2"].value, di["C3"].value],
        "cds_ytg_days": int(cds["I9"].value),
        # Vértices completos, para o CORRIGIDO (o LEGADO usa só os 19 primeiros de
        # inflação e DI e fixa os prazos pela posição na grade).
        "inflation_curve_days_full": days(cached["Dashboard"], "B", 11, 31),
        "inflation_curve_pct_full": col(dash, "C", 11, 31),
        "di_curve_days_full": days(cached["Dashboard"], "E", 11, 29),
        "cds_days": days(cached["CDS"], "C", 8, 16),
        "_celulas": {
            "first_year": "Dashboard!F3",
            "inflation_curve_pct": "Dashboard!C11:C29",
            "di_curve_pct": "Dashboard!F11:F29",
            "cds_bps": "Dashboard!I9:I17",
            "ipca_monthly": "Dashboard!I20:I31",
            "selic_monthly": "Dashboard!L20:L31",
            "inflation_params": "Inflação!C2:C3",
            "di_params": "DI!C2:C3",
            "cds_ytg_days": "CDS!I9",
            "inflation_curve_days_full": "Dashboard!B11:B31",
            "inflation_curve_pct_full": "Dashboard!C11:C31",
            "di_curve_days_full": "Dashboard!E11:E29",
            "cds_days": "CDS!C8:C16",
        },
    }


class Solution:
    """Acesso a células de uma solução do motor ``formulas``."""

    def __init__(self, sol: Any) -> None:
        self.sol = sol

    def get(self, sheet: str, ref: str) -> Any:
        key = f"'{BOOK}{sheet.upper()}'!{ref}"
        raw = self.sol[key].value
        value = raw.ravel()[0] if hasattr(raw, "ravel") else raw
        if isinstance(value, formulas.tokens.operand.XlError):
            return {"erro": str(value)}
        if isinstance(value, str):
            return value
        return float(value)

    def column(self, sheet: str, letter: str, first: int, last: int) -> list[Any]:
        return [self.get(sheet, f"{letter}{r}") for r in range(first, last + 1)]

    def row(self, sheet: str, letters: str, row: int) -> list[Any]:
        return [self.get(sheet, f"{c}{row}") for c in letters]


def rate_sheet(s: Solution, sheet: str, *, grids: bool, last_row: int) -> dict[str, Any]:
    out: dict[str, Any] = {
        "T11:T22": s.column(sheet, "T", 11, 22),
        "U11:U22": s.column(sheet, "U", 11, 22),
        "S3:S6": s.column(sheet, "S", 3, 6),
        "R26": s.get(sheet, "R26"),
        "R28:R37": s.column(sheet, "R", 28, 37),
        "S26": s.get(sheet, "S26"),
        "S27": s.get(sheet, "S27"),
        "S28:S37": s.column(sheet, "S", 28, 37),
        "E4:O4": s.row(sheet, COLS_E_TO_O, 4),
        "H6": s.get(sheet, "H6"),
        "Y3": s.get(sheet, "Y3"),
    }
    if grids:
        out["X3:X"] = s.column(sheet, "X", 3, last_row)
        out["Y3:Y"] = s.column(sheet, "Y", 3, last_row)
    return out


def cds_sheet(s: Solution, *, grids: bool) -> dict[str, Any]:
    out: dict[str, Any] = {
        "F8:F16": s.column("CDS", "F", 8, 16),
        "I9:I19": s.column("CDS", "I", 9, 19),
        "J9:J19": s.column("CDS", "J", 9, 19),
        "E4:O4": s.row("CDS", COLS_E_TO_O, 4),
    }
    if grids:
        out["M8:M"] = s.column("CDS", "M", 8, CDS_LAST_ROW)
        out["N8:N"] = s.column("CDS", "N", 8, CDS_LAST_ROW)
    return out


def dashboard(s: Solution) -> dict[str, Any]:
    return {
        "F3:P3": s.row("Dashboard", COLS_F_TO_P, 3),
        "F4:P4": s.row("Dashboard", COLS_F_TO_P, 4),
        "F5:P5": s.row("Dashboard", COLS_F_TO_P, 5),
        "F6:P6": s.row("Dashboard", COLS_F_TO_P, 6),
        "Q4": s.get("Dashboard", "Q4"),
        "Q5": s.get("Dashboard", "Q5"),
        "Q6": s.get("Dashboard", "Q6"),
    }


def outputs(s: Solution, *, grids: bool) -> dict[str, Any]:
    return {
        "Inflação": rate_sheet(s, "Inflação", grids=grids, last_row=INF_LAST_ROW),
        "DI": rate_sheet(s, "DI", grids=grids, last_row=DI_LAST_ROW),
        "CDS": cds_sheet(s, grids=grids),
        "Dashboard": dashboard(s),
    }


def scenario_inputs(sc: dict[str, Any]) -> dict[str, Any]:
    cells: dict[str, Any] = {}
    for sheet, key in (("INFLAÇÃO", "inflation_params"), ("DI", "di_params")):
        if key in sc:
            month, quarter = sc[key]
            cells[f"'{BOOK}{sheet}'!C2"] = month
            cells[f"'{BOOK}{sheet}'!C3"] = quarter
    if "cds_ytg_days" in sc:
        cells[f"'{BOOK}CDS'!I9"] = sc["cds_ytg_days"]
    return cells


def check_overrides(sol: Solution, sc: dict[str, Any]) -> None:
    """O motor ignora chave de input inválida em silêncio; aqui confirmamos que pegou."""
    for sheet, key in (("Inflação", "inflation_params"), ("DI", "di_params")):
        if key in sc:
            got = [sol.get(sheet, "C2"), sol.get(sheet, "C3")]
            assert got == sc[key], f"{sc['id']}: override de {sheet} não aplicado: {got}"
    if "cds_ytg_days" in sc:
        got_i9 = sol.get("CDS", "I9")
        assert got_i9 == sc["cds_ytg_days"], f"{sc['id']}: CDS!I9 = {got_i9}"


def dump(path: Path, data: Any) -> None:
    path.write_text(json.dumps(data, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")


def main() -> None:
    sys.setrecursionlimit(100_000)
    OUT.mkdir(parents=True, exist_ok=True)
    meta = {
        "fonte": "docs/legado/Query-Avila.xlsx",
        "fonte_sha256": sha256(XLSX),
        "motor_recalculo": f"formulas {version('formulas')}",
        "gerador": "tools/gerar_fixtures_legado.py",
    }

    inputs = read_inputs()
    dump(OUT / "inputs_planilha.json", {"_meta": meta, **inputs})
    if "--inputs-only" in sys.argv:
        return

    model = formulas.ExcelModel().loads(str(XLSX)).finish()
    base = Solution(model.calculate())
    dump(OUT / "esperado_planilha.json", {"_meta": meta, **outputs(base, grids=True)})

    scenarios = []
    for sc in SCENARIOS:
        sol = Solution(model.calculate(inputs=scenario_inputs(sc)))
        check_overrides(sol, sc)
        scenarios.append({**sc, "esperado": outputs(sol, grids=False)})
        print(f"cenário {sc['id']}: ok", flush=True)
    dump(OUT / "esperado_cenarios.json", {"_meta": meta, "cenarios": scenarios})


if __name__ == "__main__":
    main()
